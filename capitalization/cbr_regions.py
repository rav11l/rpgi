import os as _os; _os.chdir(_os.path.dirname(_os.path.abspath(__file__)))  # запуск из любой папки
import openpyxl, csv
import os; S=os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data", "sources", "cbr") + os.sep
M={"Январь":1,"Февраль":2,"Март":3,"Апрель":4,"Май":5,"Июнь":6,"Июль":7,"Август":8,"Сентябрь":9,"Октябрь":10,"Ноябрь":11,"Декабрь":12}
spec={"vol_all":("02_11_New_loans_mortgage.xlsx","в рублях"),"vol_ddu":("02_16_New_loans_scpa_mortgage.xlsx","в рублях"),
"n_all":("02_10_Quantity_mortgage.xlsx","в рублях"),"n_ddu":("02_15_Quantity_scpa_mortgage.xlsx","в рублях"),
"r_all":("02_13_Rates_mortgage.xlsx","ставка в рублях"),"r_ddu":("02_17_Rates_scpa_mortgage.xlsx","ставка в рублях"),
"t_all":("02_13_Rates_mortgage.xlsx","срок в руб"),"t_ddu":("02_17_Rates_scpa_mortgage.xlsx",None)}
data={}
for var,(f,sh) in spec.items():
    wb=openpyxl.load_workbook(S+f,read_only=True,data_only=True)
    if sh is None: sh=[s for s in wb.sheetnames if s.startswith("срок") and "руб" in s][0]
    ws=wb[sh]; rows=list(ws.iter_rows(values_only=True))
    hi=next(i for i,r in enumerate(rows) if any(isinstance(c,str) and c.split(" ")[0] in M for c in r if c))
    hdr=rows[hi]; cols={}
    for i,h in enumerate(hdr):
        if isinstance(h,str):
            p=h.split()
            if p[0] in M: cols[i]=(int(p[1]),M[p[0]])
    first=min(cols); 
    if first-1>0 and hdr[first-1] is None:
        y,m=cols[first]; cols[first-1]=(y,m-1) if m>1 else (y-1,12)
    print(var,sh,rows[1][0][:60] if rows[1][0] else "",min(cols.values()),max(cols.values()))
    for r in rows[hi+1:]:
        if not r or not isinstance(r[0],str): continue
        reg=r[0].strip()
        for i,(y,m) in cols.items():
            v=r[i]
            if isinstance(v,(int,float)): data.setdefault((reg,y,m),{})[var]=float(v)
with open("cbr_monthly.csv","w",newline="") as f:
    w=csv.writer(f); keys=list(spec); w.writerow(["region","year","month"]+keys)
    for (reg,y,m),d in sorted(data.items()):
        w.writerow([reg,y,m]+[d.get(k,"") for k in keys])
