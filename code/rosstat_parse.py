# -*- coding: utf-8 -*-
"""Разбор книг Росстата «Средние цены на первичном/вторичном рынке жилья
по субъектам и центрам субъектов РФ»: лист «по центру субъекта РФ»,
графа «Все типы квартир», I и II кварталы."""
import openpyxl, glob, re, os

def centre_sheet(wb):
    for n in wb.sheetnames:
        if "центр" in n.lower(): return wb[n]
    return wb["2"] if "2" in wb.sheetnames else None

def book_title(wb, ws):
    for r in wb[wb.sheetnames[0]].iter_rows(max_row=10, values_only=True):
        for x in r:
            if isinstance(x, str) and "рынке жилья" in x: return x
    return str(ws.cell(1,2).value or "")

def parse(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = centre_sheet(wb)
    if ws is None: return None
    title = book_title(wb, ws)
    market = "perv" if "первичн" in title.lower() else ("vtor" if "вторичн" in title.lower() else None)
    m = re.search(r"в (20\d\d)", title) or re.search(r"(20\d\d)", os.path.basename(path))
    year = int(m.group(1))
    hdr = next(r for r in range(1,12) if any("Все типы" in str(ws.cell(r,c).value or "") for c in range(1,30)))
    cols = {}
    for c in range(1,30):
        if "Все типы" in str(ws.cell(hdr,c).value or ""):
            q = str(ws.cell(hdr-1,c).value or "").strip()
            cols[q] = c
    data = {}
    for r in range(hdr+1, ws.max_row+1):
        name = str(ws.cell(r,1).value or "")
        if " - " not in name: continue
        reg = name.split(" - ",1)[1].strip()
        vals = {}
        for q, c in cols.items():
            v = ws.cell(r,c).value
            vals[q] = float(v) if isinstance(v,(int,float)) else None
        data[reg] = vals
    wb.close()
    return dict(path=os.path.basename(path), market=market, year=year, title=title[:70], data=data)

def edition_rank(fn):
    if "4kv" in fn: return 0
    if "2kv" in fn: return 1
    if "MJil" in fn or "mJil" in fn: return 2
    if "1kv" in fn: return 4
    return 3
